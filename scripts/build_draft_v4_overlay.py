"""
draft_v4_overlay.docx — KCC2026_심사용_원본.docx 위에 본문만 얹기

핵심 변경 (vs build_draft_v3_docx.py):
  - rFonts 강제 지정 제거: 심사용 원본의 Normal 스타일(바탕) + docDefaults(Times New Roman)
    가 자동으로 적용되어 sample1.pdf와 동일한 폰트로 출력됨
  - Figure 1(파이프라인)은 본문에 삽입하지 않음 — 사용자가 별도 제작
    (output/paper_figures/_diagram_to_make/ 에 후보 이미지 분리)
  - sectPr 2개(1-col 제목/요약 + 2-col 본문) 그대로 유지
  - 본문 콘텐츠/표 데이터는 v3와 동일 (이미 검증된 수치)
"""
import os, shutil
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = '/Users/honghwasu/Desktop/research_project/output/논문양식'
SRC  = f'{BASE}/KCC2026_심사용_원본.docx'
DST  = f'{BASE}/draft_v4_overlay.docx'

FIGS_DIR = '/Users/honghwasu/Desktop/research_project/output/paper_figures'
FIG2 = f'{FIGS_DIR}/fig2_3d_visualization_v5.png'

# 폰트 크기 — 심사용 원본/sample1.pdf 실측 매칭 (본문 10pt 기준)
SIZE_TITLE     = 16
SIZE_TITLE_EN  = 14
SIZE_SECTION   = 12
SIZE_SUBSEC    = 10
SIZE_BODY      = 10
SIZE_SUMMARY   = 10
SIZE_TABLE     = 9
SIZE_REF       = 9
SIZE_CAP       = 9

# ============================================================
# Build base
# ============================================================
shutil.copy(SRC, DST)
doc = Document(DST)
body = doc.element.body

# sectPr 보존하며 본문 비우기
section_break_paras, other, final_sectPr = [], [], None
for child in list(body):
    if child.tag == qn('w:sectPr'):
        final_sectPr = child
    elif child.tag == qn('w:p'):
        pPr = child.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
            section_break_paras.append(child)
        else:
            other.append(child)
    elif child.tag == qn('w:tbl'):
        other.append(child)

for c in other:
    body.remove(c)

# section break paragraph 안의 run 제거 (sectPr만 남김)
for sp in section_break_paras:
    for r in list(sp):
        if r.tag != qn('w:pPr'):
            sp.remove(r)

assert len(section_break_paras) == 1 and final_sectPr is not None, \
    f"Section 구조 이상: break_paras={len(section_break_paras)}, final={final_sectPr is not None}"
sec_break = section_break_paras[0]

# ============================================================
# Helpers — 폰트 강제 지정 제거 (Normal 스타일 inherit)
# ============================================================
def make_para(text, size, bold=False, italic=False,
              align=None, indent_mm=0, space_after_pt=3, line_spacing=None):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    if align is not None:
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), {
            WD_ALIGN_PARAGRAPH.CENTER: 'center',
            WD_ALIGN_PARAGRAPH.LEFT: 'left',
            WD_ALIGN_PARAGRAPH.JUSTIFY: 'both',
            WD_ALIGN_PARAGRAPH.RIGHT: 'right'
        }[align])
        pPr.append(jc)
    if indent_mm:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), str(int(indent_mm * 56.69)))
        pPr.append(ind)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:after'), str(int(space_after_pt * 20)))
    if line_spacing:
        spacing.set(qn('w:line'), str(int(line_spacing * 240)))
        spacing.set(qn('w:lineRule'), 'auto')
    pPr.append(spacing)
    p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size * 2))); rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement('w:b'))
        rPr.append(OxmlElement('w:bCs'))
    if italic:
        rPr.append(OxmlElement('w:i'))
    # NOTE: rFonts 강제 지정 제거 → docDefaults(영문 Times New Roman / 한글 바탕) 적용
    r.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    p.append(r)
    return p

def heading(text, level='section'):
    size = {'title': SIZE_TITLE, 'title_en': SIZE_TITLE_EN,
            'section': SIZE_SECTION, 'sub': SIZE_SUBSEC}[level]
    return make_para(text, size=size, bold=True, space_after_pt=1)

def table_elem(headers, rows, col_widths_mm):
    tbl = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    st = OxmlElement('w:tblStyle'); st.set(qn('w:val'), 'TableGrid'); tblPr.append(st)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(sum(col_widths_mm) * 56.69)))
    tblW.set(qn('w:type'), 'dxa'); tblPr.append(tblW)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), '000000')
        borders.append(b)
    tblPr.append(borders)
    tbl.append(tblPr)
    tblGrid = OxmlElement('w:tblGrid')
    for w_mm in col_widths_mm:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(w_mm * 56.69))); tblGrid.append(gc)
    tbl.append(tblGrid)

    def cell(txt, bold=False):
        tc = OxmlElement('w:tc')
        tcPr = OxmlElement('w:tcPr')
        tcW = OxmlElement('w:tcW'); tcW.set(qn('w:type'), 'auto'); tcW.set(qn('w:w'), '0'); tcPr.append(tcW)
        vA = OxmlElement('w:vAlign'); vA.set(qn('w:val'), 'center'); tcPr.append(vA)
        tc.append(tcPr)
        p = make_para(str(txt), size=SIZE_TABLE, bold=bold,
                      align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=0)
        tc.append(p)
        return tc

    tr = OxmlElement('w:tr')
    for h in headers:
        tr.append(cell(h, bold=True))
    tbl.append(tr)
    for row in rows:
        tr = OxmlElement('w:tr')
        for v in row:
            tr.append(cell(v))
        tbl.append(tr)
    return tbl

def insert_before(anchor, *els):
    for el in els:
        anchor.addprevious(el)

# ============================================================
# SECTION 1 (1단): 제목 + 영문제목 + 요약
# ============================================================
s1 = []
s1.append(make_para(
    'GeoCalib + Depth Pro 기반 부동산 매물 사진의 실내 벽-벽 거리 측정 및 실패 케이스 수동 보완 프레임워크',
    size=SIZE_TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=3, line_spacing=1.15))
s1.append(make_para(
    'A Framework for Indoor Wall-to-Wall Distance Measurement in Real Estate Photos '
    'using GeoCalib and Depth Pro, with Manual Refinement for RANSAC Failures',
    size=SIZE_TITLE_EN, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=12))

s1.append(heading('요   약', level='sub'))
s1.append(make_para(
    '부동산 매물 사진은 광각 렌즈로 촬영되어 실제보다 공간이 넓어 보이며, 플랫폼이 EXIF 정보를 삭제하여 '
    '소비자가 실내 치수를 사전에 확인하기 어렵다. 본 논문은 EXIF가 없는 단일 실내 사진에서 초점거리를 '
    'GeoCalib으로 추정하고, Depth Pro로 metric depth를 생성한 뒤, RANSAC 평면 피팅으로 평행 벽 쌍을 '
    '탐지하여 벽-벽 거리를 자동 측정하는 프레임워크를 제안한다. 5개 단지 아파트 81장의 평면도 기준 거리를 '
    'GT로 사용하여 정량 검증한 결과, 측정 실패(RANSAC 실패 2장 + 극단 오차 ≥80% 6장) 8장을 제외한 '
    '73장에서 전체 MAPE 18.2%, 중앙값 15.2%, 38.4%(28장)가 10% 이하 오차를 기록하였으며, '
    '양호 조건 25장(전체의 30.9%)에서는 76%(19장)가 10% 이하 오차(중앙값 5.3%, MAE 28.8 cm)를 기록하였다. '
    'GeoCalib 추정 초점거리 분포의 69.8%(44/63)가 500 px 미만으로 광각 렌즈 사용을 시사하였다. '
    '측정 실패 케이스를 제외한 paired 비교(N=55)에서 GeoCalib focal이 Depth Pro 자체 focal 추정 대비 '
    "평균 1.66%p 낮은 오차(paired t-test p=0.039, Cohen's d=-0.29)를 보였고, 광각 subset(N=39)에서는 "
    '효과가 2.47%p (p=0.018, d=-0.40)로 강화되었다. 또한 자동 RANSAC이 실패하거나 '
    '큰 오차를 보인 케이스에 대해 사용자가 벽 영역을 수동 지정하면 장애물 17장(paired)에서 평균 오차가 '
    '23.5%→19.2%, RANSAC 잘못 피팅 5장에서 53.2%→25.0%로 개선되어(최대 개선 69.9%p), 자동-수동 '
    '하이브리드 측정의 실용적 가능성을 시사하는 탐색적 결과를 얻었다(장애물 Wilcoxon p=0.057).',
    size=SIZE_SUMMARY, indent_mm=4, space_after_pt=6, line_spacing=1.0))

insert_before(sec_break, *s1)

# ============================================================
# SECTION 2 (2단): 본문 + 참고문헌
# ============================================================
s2 = []

# 1. 서론
s2.append(heading('1. 서 론', level='section'))
for t in [
    '부동산 매물 정보에서 전용면적은 공개되지만 개별 방의 가로·세로 치수는 제공되지 않는 경우가 많다. '
    '소비자가 공간 크기를 사전에 판단할 수 있는 거의 유일한 시각 정보는 매물 사진이나, 해당 사진은 대부분 '
    '광각 렌즈로 촬영되어 실제보다 공간이 넓어 보이는 왜곡을 발생시킨다. 주요 부동산 플랫폼은 이미지 '
    '업로드 시 EXIF 메타데이터를 삭제하므로 소비자는 촬영에 사용된 초점거리나 렌즈 종류를 확인할 수 없다.',
    '본 연구의 장기 목표는 원룸·오피스텔 등 치수 정보가 부재한 소형 매물 사진에서 실내 공간 거리를 '
    '추정하는 것이다. 그러나 소형 매물은 평면도 GT 확보가 어려워 정량 검증이 제한되므로, 본 논문에서는 '
    '평면도 치수를 확보할 수 있는 아파트 사진을 대상으로 파이프라인을 먼저 구축·검증한다.',
    '본 논문은 EXIF가 없는 단일 실내 사진에서 초점거리를 GeoCalib [1]으로 추정하고, Depth Pro [2]로 '
    'metric depth를 생성한 뒤, 3D 역투영과 RANSAC 평면 피팅을 통해 평행 벽 쌍을 자동 탐지하여 벽-벽 '
    '거리를 측정하는 프레임워크를 제안한다(그림 1). 5개 단지 아파트 81장의 평면도 GT에 대해 정량 평가, '
    '조건별 오차 분포, 초점거리 추정 방식 비교, 자동 실패 시 수동 보완 실험까지 분석한다.',
]:
    s2.append(make_para(t, size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(make_para(
    '본 논문의 주요 기여는 (1) EXIF가 없는 부동산 매물 사진에서 GeoCalib · Depth Pro · RANSAC을 '
    '결합한 벽-벽 거리 측정 파이프라인 제안, (2) 5개 단지 81장에 대한 평면도 GT 기반 정량 검증(측정 '
    '실패 8장 제외 73장 전체 MAPE 18.2%, 양호 조건 25장 76%가 10% 이하 오차), (3) Depth Pro 자체 '
    "focal 추정과의 paired 비교에서 GeoCalib의 통계적 우위 확인(전체 N=55 p=0.039, 광각 N=39 p=0.018), "
    '(4) RANSAC 실패 케이스에 대한 수동 보완 예비 실험(최대 69.9%p 개선, 장애물 Wilcoxon p=0.057)으로 '
    '자동-수동 하이브리드 인터페이스의 예비 근거 제시이다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

# 2. 관련 연구
s2.append(heading('2. 관련 연구', level='section'))
s2.append(make_para(
    '카메라 캘리브레이션은 전통적 체커보드 기반 [9]에서 딥러닝 단일 이미지 접근 [6]으로 발전했고, '
    'GeoCalib [1]은 기하 최적화와 신경망 추정을 결합한다. 단안 metric depth는 Metric3D v2 [3], '
    'UniDepth [4], Depth Anything V2 [5]를 거쳐, Depth Pro [2]가 카메라 파라미터 없이 고해상도 '
    'metric depth를 수 초 내 생성한다. 부동산·실내 응용으로 Rent3D [7]는 사진-평면도 결합을, '
    'Zillow Indoor Dataset [8]은 파노라마-평면도 데이터셋을 제공한다. 국내에서 권혁찬 등 [10]은 '
    'GeoCalib을 매물 사진 왜곡 판별(Propix)에 적용하였으나 정량 치수 측정은 다루지 않는다. '
    '본 연구는 GeoCalib + Depth Pro + RANSAC 결합으로 mm 단위 벽-벽 거리를 측정·검증한다는 점에서 차별화된다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

# 3. 제안 프레임워크
s2.append(heading('3. 제안 프레임워크', level='section'))
s2.append(make_para(
    '제안 프레임워크는 EXIF가 없는 단일 실내 사진으로부터 벽-벽 거리를 자동 측정하는 6단계 자동 경로와, '
    'RANSAC이 실패하거나 잘못 피팅하는 케이스에 대한 사용자 수동 보완 경로로 구성된다. 전체 흐름은 '
    '그림 1에 블록도로 도시하였다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

# Figure 1 placeholder (사용자가 직접 제작)
s2.append(make_para('[그림 1. 전체 파이프라인 흐름도 — 별도 제작 예정]',
                    size=SIZE_CAP, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=6))

s2.append(heading('3.1 GeoCalib + Depth Pro (초점거리·깊이)', level='sub'))
s2.append(make_para(
    '입력 사진 I를 GeoCalib [1]에 입력해 초점거리 f_px, 롤·피치 각, 방사 왜곡 계수 k1·k2를 추정한다 '
    '(Perspective Field 신경망 + Levenberg-Marquardt 최적화, distorted weight). 추정된 f_px를 Depth Pro [2]에 '
    '입력해 동일 해상도 metric depth map D를 얻는다. Depth Pro는 canonical inverse depth를 (W / f_px)로 '
    '스케일링해 실제 미터 단위 깊이를 생성하며, GeoCalib focal 주입이 자체 추정 대비 광각에서 더 정확한 '
    'depth를 산출함은 4.4절에서 확인한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('3.2 3D 역투영 + RANSAC 평면 피팅 + 수동 보완', level='sub'))
s2.append(make_para(
    '픽셀 (u, v)의 3D 좌표를 X = (u − cx) · Z / f_px, Y = (v − cy) · Z / f_px, Z = D(u, v)로 복원하며 '
    '유효 범위를 0.5 < Z < 10 m로 제한한다. 점군에 대해 τ = 5 cm, iter = 1000 RANSAC을 반복해 최대 '
    '6개 평면을 순차 추출하고, |ny| ≤ 0.7인 평면만 수직 벽 후보로 보존한다. 이후 |n_i · n_j| > 0.8 '
    '평행 쌍 중 inlier 합 최대 쌍을 최종 선택하여 벽-벽 거리 dist = |d_i ± d_j|를 산출한다. 본 임계값은 '
    'Depth Pro의 per-pixel noise를 허용하는 설정이며 엄격한 ablation은 향후 과제이다. 자동 RANSAC 실패 '
    '시에는 사용자가 두 벽의 40 × 40 px 패치를 클릭해 해당 영역 depth 중앙값으로 거리를 직접 산출하는 '
    '수동 보완 경로를 병행하며, 그 정량 효과는 4.6절에서 보고한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(make_para('그림 2. 3D 점군+RANSAC 벽 평면+거리선 (ID 029 거실, GT 4,340 mm, 측정 4,355 mm, 오차 0.35%).',
                    size=SIZE_CAP, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=6))

# 4. 실험 및 결과
s2.append(heading('4. 실험 및 결과', level='section'))

s2.append(heading('4.1 데이터셋 및 GT 구성', level='sub'))
s2.append(make_para(
    '올림픽파크포레온, 개포자이프레지던스, 래미안원베일리, 헬리오시티, 마포래미안푸르지오의 5개 단지에서 '
    '실내 매물 사진 총 81장을 수집하였다. GT는 네이버 부동산 평면도의 벽 중심선 치수에서 벽 두께 '
    '160 mm를 차감한 마감면 간 거리를 사용한다(단지별 벽 두께 편차 ±40 mm로 GT에 약 1~2% 내재 불확실성). '
    '81장 중 RANSAC 완전 실패 2장(ID 018, 034) 및 측정 결과가 GT와 80% 이상 괴리된 극단 오류 6장'
    '(ID 014, 028, 030, 031, 063, M8)을 제외한 73장을 분석 대상으로 한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('4.2 광각 환경 검증', level='sub'))
s2.append(make_para(
    'GeoCalib 추정 초점거리(N=63) 평균 490 px, 중앙값 446 px이며 69.8%(44/63)가 500 px 미만으로, '
    '수집된 매물 사진 다수가 광각 렌즈로 촬영되었을 가능성을 정량적으로 시사한다. 4.4절 paired 비교에는 '
    'Depth Pro 자체 focal 추정이 성공한 61장(광각 43, 일반 18) 중 측정 오류를 제외한 55장을 사용한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('4.3 조건별 측정 오차', level='sub'))
s2.append(make_para('표 1. 조건별 측정 오차 통계 (MAE·RMSE: mm, MAPE·중앙값: %). 분석 대상 73장.',
                    size=SIZE_CAP, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=2))
s2.append(table_elem(
    headers=['조건', 'N', 'MAE', 'RMSE', 'MAPE', '중앙값', '≤10%'],
    rows=[
        ['전체 유효',         '73', '643',  '948',  '18.2', '15.2', '38.4%'],
        ['양호(final_clean)', '25', '288',  '390',  '7.0',  '5.3',  '76.0%'],
        ['both_clear',        '48', '489',  '752',  '13.4', '10.6', '50.0%'],
        ['one_unclear',       '25', '937',  '1241', '27.3', '23.0', '16.0%'],
        ['obstacle=yes',      '20', '818',  '980',  '23.7', '22.5', '15.0%'],
    ],
    col_widths_mm=[28, 6, 11, 12, 11, 11, 11]))
s2.append(make_para(
    '양호 조건(final_clean)은 wall_visible=both_clear, obstacle=none, 발코니 유리문 없음을 모두 만족하는 '
    '25장(전체의 30.9%)이다. 조건 subset은 상호 배타적이지 않다. 분석 대상 73장 전체에서 38.4%(28장)가 '
    '10% 이하 오차를 달성하며, 양호 조건 선별 시 76%로 상승한다. 선별 비율이 30.9%에 불과하므로 일반 '
    '매물 사진 전반 적용 시 실용 성공률은 38%대로 추정된다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('4.4 초점거리 추정 방식 비교', level='sub'))
s2.append(make_para('표 2. GeoCalib focal (A) vs Depth Pro 자체 focal (B) paired 비교 '
                    "(MAPE: %, Δ: %p, d: Cohen's d). 측정 실패(err > 80%) 케이스 제외 N=55.",
                    size=SIZE_CAP, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=2))
s2.append(table_elem(
    headers=['그룹', 'N', 'A', 'B', 'Δ (95% CI)', 'p', 'd'],
    rows=[
        ['전체',        '55', '18.04', '19.70', '−1.66 (−3.24, −0.08)',  '0.039', '−0.285'],
        ['광각 <500px', '39', '19.03', '21.49', '−2.47 (−4.49, −0.45)',  '0.018', '−0.396'],
        ['일반 ≥500px', '16', '15.62', '15.33', '+0.29 (−2.02, +2.60)',  '0.790', '+0.068'],
    ],
    col_widths_mm=[16, 6, 9, 9, 30, 9, 11]))
s2.append(make_para(
    '측정 자체가 실패한 케이스(err > 80% on A or B)는 focal 정확도 비교의 본질과 무관하므로 사전 '
    '제외 기준으로 N=61에서 N=55로 축소하였다. 전체(N=55) 수준에서 GeoCalib focal이 Depth Pro 자체 '
    "추정 대비 1.66%p 낮은 오차(paired t-test p=0.039, Cohen's d=-0.285)로 유의하게 우수하였으며, "
    '광각(<500 px, N=39) 서브셋에서는 2.47%p 낮은 오차(p=0.018, d=-0.396)로 더 강한 효과를 보였다. '
    '세 그룹에 Bonferroni 보정(α=0.0167) 적용 시 광각 p=0.018과 전체 p=0.039가 모두 α를 근소하게 '
    '상회하여 엄격히는 유의하지 않음(marginal)을 투명히 명시한다. 다만 광각이 사전 설정된 주 관심 '
    '그룹임을 감안한 단일 검정 관점에서는 p=0.018이 유의 수준을 만족한다. 일반 화각에서는 '
    '두 방식의 차이가 유의하지 않아, GeoCalib focal 결합의 이점이 광각 조건에서 집중됨을 시사한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('4.5 실패 모드 분석', level='sub'))
s2.append(make_para(
    '자동 RANSAC의 주요 실패 모드는 (1) 장애물(붙박이장·가구·주방 아일랜드가 벽 평면을 대체), '
    '(2) 발코니 유리문(유리 너머 depth 누락), (3) 한쪽 벽 미보임(수평 점군 불충분)의 세 가지이다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

s2.append(heading('4.6 실패 케이스 수동 보완 측정', level='sub'))
s2.append(make_para(
    '사용자가 두 벽의 패치(40×40 px)를 클릭하면 depth 중앙값으로 거리를 산출하는 수동 보완 실험을 수행하였다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))
s2.append(make_para('표 3. 자동 RANSAC vs 수동 클릭 기반 측정 비교 (MAPE 단위: %, 개선 단위: %p).',
                    size=SIZE_CAP, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=2))
s2.append(table_elem(
    headers=['Subset', 'N', '자동%', '수동%', 'Δ', 'p', '개선'],
    rows=[
        ['장애물',         '17', '23.48', '19.22', '−4.26',  '0.057', '13/17'],
        ['RANSAC 오피팅', '5',  '53.18', '25.02', '−28.16', '0.188', '4/5'],
    ],
    col_widths_mm=[16, 5, 10, 10, 9, 8, 8]))
s2.append(make_para(
    '장애물 유형별 개선 효과는 편차가 컸다: 가구(식탁·의자·소파) 5장 평균 +16.3%p (ID 60: 65.1%→1.4%), '
    '붙박이장 5장 평균 +2.3%p, 주방 아일랜드 6장(반사성 표면) 평균 −3.3%p. RANSAC이 창문·유리문·코너와 '
    '혼동한 실패 5장에서는 평균 28.2%p 개선되었다(ID 14: 96.9%→27.0%, ID 51: 71.3%→5.2%). 장애물 전체 '
    'Wilcoxon p=0.057로 유의성 경계이며 일부 코너·벽 잘림 케이스에서는 수동이 오차를 악화시키는 사용자 '
    '의존성이 존재한다. 평면도 GT 없는 원룸·고성 사진 6장 예비 적용에서도 자동 측정이 합리적 수치를 '
    '산출하였다(원룸 거실 4.36 m, 작은방 3.04 m — 일반 한국 매물 범위 내).',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

# 5. 결론
s2.append(heading('5. 결론 및 한계', level='section'))
s2.append(make_para(
    '본 논문은 EXIF 없는 매물 사진에서 GeoCalib + Depth Pro + RANSAC 결합 프레임워크를 제안하여 5개 단지 '
    '81장(측정 실패 제외 73장)에서 38.4%(양호 조건 76%)가 10% 이하 오차를 달성하였고, 광각 GeoCalib '
    'focal 우위(p=0.018)와 수동 보완 최대 69.9%p 개선을 관찰하였다. 한계는 단일 depth 모델 검증과 '
    'RANSAC ablation 부재이며, 향후 원룸·오피스텔 GT 확보, Metric3D v2/UniDepth 교차 검증, 자동-수동 '
    '하이브리드 인터페이스 설계를 계획한다.',
    size=SIZE_BODY, indent_mm=4, space_after_pt=2, line_spacing=1.0))

# 참고문헌
s2.append(heading('참고 문헌', level='sub'))
for ref in [
    '[1] A. Veicht et al., "GeoCalib," ECCV 2024.',
    '[2] A. Bochkovskii et al., "Depth Pro," ICLR 2025.',
    '[3] M. Hu et al., "Metric3D v2," TPAMI 2024.',
    '[4] L. Piccinelli et al., "UniDepth," CVPR 2024.',
    '[5] L. Yang et al., "Depth Anything V2," NeurIPS 2024.',
    '[6] L. Jin et al., "Perspective Fields for Single Image Camera Calibration," CVPR 2023.',
    '[7] C. Liu et al., "Rent3D," CVPR 2015.',
    '[8] S. Cruz et al., "Zillow Indoor Dataset," CVPR 2021.',
    '[9] Z. Zhang, "Flexible Camera Calibration," IEEE PAMI 22(11), 2000.',
    '[10] 권혁찬 외, "AI 부동산 왜곡 탐지·보정," 한국정보기술학회, 2024.',
]:
    s2.append(make_para(ref, size=SIZE_REF, indent_mm=3, space_after_pt=0, line_spacing=0.9))

insert_before(final_sectPr, *s2)
doc.save(DST)
print(f'[step1] text saved: {DST}')

# ============================================================
# 그림 2 삽입 (Figure 1은 사용자가 직접 제작)
# ============================================================
doc = Document(DST)
fig_map = [
    ('그림 2.', FIG2, 80),
]
for cap_prefix, path, width_mm in fig_map:
    if not os.path.exists(path):
        print(f'  MISSING {path}')
        continue
    inserted = False
    for p in doc.paragraphs:
        if p.text.strip().startswith(cap_prefix):
            new_p = p.insert_paragraph_before()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.add_run().add_picture(path, width=Mm(width_mm))
            inserted = True
            print(f'  inserted {cap_prefix} ← {os.path.basename(path)}')
            break
    if not inserted:
        print(f'  placeholder not found for {cap_prefix}')

doc.save(DST)
print(f'\n[DONE] {DST}  size={os.path.getsize(DST):,} bytes')
